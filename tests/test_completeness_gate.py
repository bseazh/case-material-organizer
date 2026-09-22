from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "lawyerbuddy-sorting" / "scripts"
sys.path.insert(0, str(SCRIPTS))

from completeness import require_completeness, validate_completeness  # noqa: E402


def complete_plan(unit_type: str = "page", expected: int = 10) -> dict:
    return {
        "items": [{"material_id": "MAT-0001", "original_name": "测试材料.pdf"}],
        "reading_coverage": {
            "materials": [{
                "material_id": "MAT-0001",
                "unit_type": unit_type,
                "units_expected": expected,
                "units_completed": expected,
                "source_units": [str(index) for index in range(1, expected + 1)],
                "completed_units": [str(index) for index in range(1, expected + 1)],
                "status": "complete",
                "coverage_basis": "原文件",
                "no_relevant_fact_reason": "",
                "notes": "",
            }],
            "review_rounds": {
                "full_extraction": {
                    "completed": True,
                    "reviewed_material_ids": ["MAT-0001"],
                    "reviewed_fact_ids": ["FACT-001"],
                },
                "cross_material_review": {
                    "completed": True,
                    "reviewed_material_ids": ["MAT-0001"],
                    "reviewed_fact_ids": ["FACT-001"],
                },
                "legal_fact_review": {
                    "completed": True,
                    "reviewed_material_ids": ["MAT-0001"],
                    "reviewed_fact_ids": ["FACT-001"],
                },
            },
        },
        "fact_inventory": [{
            "fact_id": "FACT-001",
            "statement": "材料末尾记载了与履行期限有关的事实。",
            "record_nature": "书面约定",
            "source_material_ids": ["MAT-0001"],
            "source_locations": ["第10页"],
            "legal_relevance": ["履行期限"],
        }],
        "fact_disposition": {
            "report_body": ["FACT-001"],
            "timeline": [],
            "background": [],
            "pending_confirmation": [],
        },
        "legal_fact_map": [{
            "legal_element": "履行期限",
            "fact_ids": ["FACT-001"],
            "evidence_status": "材料已有记载",
            "issues": "",
        }],
    }


class CompletenessGateTest(unittest.TestCase):
    def assert_unit_gap_is_blocked(self, unit_type: str, expected: int, completed: int) -> None:
        plan = complete_plan(unit_type, expected)
        record = plan["reading_coverage"]["materials"][0]
        record["units_completed"] = completed
        record["completed_units"] = record["completed_units"][:completed]
        record["status"] = "partial"
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertLess(result.unit_coverage_rate, 1.0)
        self.assertTrue(any("尚未完整读取" in error for error in result.errors))
        with self.assertRaisesRegex(ValueError, "阅读单元覆盖率未达到 100%"):
            require_completeness(plan)

    def test_key_fact_on_last_pdf_page_cannot_be_skipped(self) -> None:
        self.assert_unit_gap_is_blocked("page", expected=10, completed=9)

    def test_key_fact_on_second_worksheet_cannot_be_skipped(self) -> None:
        self.assert_unit_gap_is_blocked("worksheet", expected=2, completed=1)

    def test_key_fact_in_final_transcript_segment_cannot_be_skipped(self) -> None:
        self.assert_unit_gap_is_blocked("transcript_segment", expected=8, completed=7)

    def test_every_fact_must_have_exactly_one_disposition(self) -> None:
        plan = complete_plan()
        plan["fact_disposition"]["timeline"].append("FACT-001")
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("多个主要去向" in error for error in result.errors))

    def test_legal_relevant_fact_must_be_mapped(self) -> None:
        plan = complete_plan()
        plan["legal_fact_map"] = []
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("尚未映射法律要素" in error for error in result.errors))

    def test_completed_count_cannot_replace_unit_details(self) -> None:
        plan = complete_plan(expected=10)
        plan["reading_coverage"]["materials"][0]["completed_units"] = ["1-10"]
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("明细与数量不一致" in error for error in result.errors))

    def test_expected_count_cannot_be_smaller_than_source_manifest(self) -> None:
        plan = complete_plan(expected=1)
        record = plan["reading_coverage"]["materials"][0]
        record["source_units"] = [str(index) for index in range(1, 11)]
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("来源单元清单与应读数量不一致" in error for error in result.errors))

    def test_complete_status_cannot_substitute_extra_or_missing_units(self) -> None:
        plan = complete_plan(expected=10)
        record = plan["reading_coverage"]["materials"][0]
        record["completed_units"] = [str(index) for index in range(2, 12)]
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("来源清单之外" in error for error in result.errors))

    def test_material_without_fact_requires_a_reason(self) -> None:
        plan = complete_plan()
        plan["items"].append({"material_id": "MAT-0002", "original_name": "背景材料.pdf"})
        plan["reading_coverage"]["materials"].append({
            "material_id": "MAT-0002",
            "unit_type": "page",
            "units_expected": 1,
            "units_completed": 1,
            "source_units": ["1"],
            "completed_units": ["1"],
            "status": "complete",
            "coverage_basis": "原文件",
            "no_relevant_fact_reason": "",
        })
        for review in plan["reading_coverage"]["review_rounds"].values():
            review["reviewed_material_ids"].append("MAT-0002")
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("未提取事实" in error for error in result.errors))
        plan["reading_coverage"]["materials"][1]["no_relevant_fact_reason"] = "仅为无实质内容的封面。"
        self.assertTrue(validate_completeness(plan).passed)

    def test_review_round_scope_cannot_be_empty(self) -> None:
        plan = complete_plan()
        plan["reading_coverage"]["review_rounds"]["full_extraction"]["reviewed_fact_ids"] = []
        plan["reading_coverage"]["review_rounds"]["cross_material_review"]["reviewed_material_ids"] = []
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("完整提取复核未覆盖全部事实" in error for error in result.errors))
        self.assertTrue(any("跨材料复核未覆盖全部材料" in error for error in result.errors))

    def test_legal_map_requires_evidence_status(self) -> None:
        plan = complete_plan()
        plan["legal_fact_map"][0]["evidence_status"] = ""
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("缺少证据状态" in error for error in result.errors))

    def test_all_three_review_rounds_are_required(self) -> None:
        plan = complete_plan()
        plan["reading_coverage"]["review_rounds"]["cross_material_review"]["completed"] = False
        result = validate_completeness(plan)
        self.assertFalse(result.passed)
        self.assertTrue(any("cross_material_review" in error for error in result.errors))

    def test_complete_coverage_passes(self) -> None:
        result = require_completeness(copy.deepcopy(complete_plan()))
        self.assertTrue(result.passed)
        self.assertEqual(result.material_coverage_rate, 1.0)
        self.assertEqual(result.unit_coverage_rate, 1.0)
        self.assertEqual(result.fact_disposition_rate, 1.0)

    def test_report_builder_enforces_gate_and_detailed_sections(self) -> None:
        source = (SCRIPTS / "build_report.py").read_text(encoding="utf-8")
        timeline_source = (SCRIPTS / "build_timeline.py").read_text(encoding="utf-8")
        self.assertIn("require_analysis_readiness(plan", source)
        self.assertIn("require_analysis_readiness(plan", timeline_source)
        self.assertIn("2.6 重要事实完整梳理", source)
        self.assertIn("2.7 法律事实及要素对应", source)

    def test_extraction_records_page_sheet_and_segment_units(self) -> None:
        source = (SCRIPTS / "extract_content.py").read_text(encoding="utf-8")
        self.assertIn("pdf_page_count", source)
        self.assertIn('"reading_unit_type"', source)
        self.assertIn('"units_expected"', source)
        self.assertIn('"completed_units"', source)

    def test_extractor_reads_second_worksheet(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "source"
            source.mkdir()
            workbook = Workbook()
            workbook.active.title = "付款记录"
            workbook.active["A1"] = "已付款100000元"
            second = workbook.create_sheet("退款差额")
            second["A1"] = "退款20000元，尚欠30000元"
            workbook.save(source / "付款明细.xlsx")
            inventory = {
                "source_folder": str(source),
                "files": [{"original_relative_path": "付款明细.xlsx", "extension": ".xlsx"}],
            }
            inventory_file = temp / "inventory.json"
            inventory_file.write_text(json.dumps(inventory, ensure_ascii=False), encoding="utf-8")
            output_dir = temp / "output"
            subprocess.run(
                [sys.executable, str(SCRIPTS / "extract_content.py"), str(inventory_file), "--out-dir", str(output_dir)],
                check=True,
                capture_output=True,
                text=True,
            )
            extraction = json.loads((output_dir / "extractions.json").read_text(encoding="utf-8"))["files"][0]
            extracted_text = (output_dir / extraction["text_file"]).read_text(encoding="utf-8")
            self.assertEqual(extraction["source_units"], ["付款记录", "退款差额"])
            self.assertEqual(extraction["units_expected"], 2)
            self.assertIn("尚欠30000元", extracted_text)

    def test_extractor_reads_final_long_document_segment(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "source"
            source.mkdir()
            document = Document()
            document.add_paragraph("前段沟通。" * 2500)
            document.add_paragraph("最后主张解除合同并退还全部款项。")
            document.save(source / "沟通逐字稿.docx")
            inventory = {
                "source_folder": str(source),
                "files": [{"original_relative_path": "沟通逐字稿.docx", "extension": ".docx"}],
            }
            inventory_file = temp / "inventory.json"
            inventory_file.write_text(json.dumps(inventory, ensure_ascii=False), encoding="utf-8")
            output_dir = temp / "output"
            subprocess.run(
                [sys.executable, str(SCRIPTS / "extract_content.py"), str(inventory_file), "--out-dir", str(output_dir)],
                check=True,
                capture_output=True,
                text=True,
            )
            extraction = json.loads((output_dir / "extractions.json").read_text(encoding="utf-8"))["files"][0]
            extracted_text = (output_dir / extraction["text_file"]).read_text(encoding="utf-8")
            self.assertGreaterEqual(extraction["units_expected"], 2)
            self.assertEqual(extraction["units_expected"], len(extraction["source_units"]))
            self.assertIn("最后主张解除合同并退还全部款项", extracted_text)

    def test_incomplete_plan_cannot_generate_final_timeline(self) -> None:
        plan = complete_plan(expected=10)
        record = plan["reading_coverage"]["materials"][0]
        record["units_completed"] = 9
        record["completed_units"] = record["completed_units"][:9]
        record["status"] = "partial"
        plan.update({"confirmed": True, "media_check": {}, "events": [], "case_summary": {}})
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            plan_file = temp / "plan.json"
            output = temp / "timeline.html"
            plan_file.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "build_timeline.py"), str(plan_file), "--out", str(output)],
                capture_output=True,
                text=True,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("阅读单元覆盖率未达到 100%", result.stderr)
            self.assertFalse(output.exists())

    def test_complete_plan_generates_detailed_word_report(self) -> None:
        plan = complete_plan()
        plan.update({
            "confirmed": True,
            "media_check": {},
            "case_folder": {
                "sequence": "1",
                "plaintiff_short_name": "张三",
                "defendant_short_name": "李四",
                "cause_of_action": "追索劳动报酬纠纷",
            },
            "cause_of_action_review": {
                "status": "confirmed",
                "primary_cause": "追索劳动报酬纠纷",
                "level": 4,
                "hierarchy": [
                    "劳动争议、人事争议、新就业形态用工纠纷",
                    "劳动争议",
                    "劳动合同纠纷",
                    "追索劳动报酬纠纷",
                ],
                "confirmed_by_user": True,
            },
            "case_folder_name": "1-张三VS李四-追索劳动报酬纠纷",
            "case_summary": {
                "起因": "材料记载双方存在劳动关系。",
                "过程": "材料记载工资支付经过。",
                "争议": "工资金额尚待核对。",
                "现状": "尚在核对材料。",
                "缺口": "需补充完整工资记录。",
            },
            "entities": [],
            "events": [],
        })
        plan["items"][0].update({
            "target_category": "002 基础资料",
            "target_subcategory": "01 劳动合同与入职",
            "proposed_name": "240101-劳动合同.pdf",
        })
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            plan["result_folder"] = str(temp / plan["case_folder_name"])
            plan_file = temp / "plan.json"
            output = temp / "追索劳动报酬纠纷案件梳理报告.docx"
            plan_file.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
            subprocess.run(
                [sys.executable, str(SCRIPTS / "build_report.py"), str(plan_file), "--out", str(output)],
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertTrue(output.is_file())
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("2.6 重要事实完整梳理", text)
            self.assertIn("2.7 法律事实及要素对应", text)


if __name__ == "__main__":
    unittest.main()

#!/usr/bin/env python3
"""Extract searchable text from inventoried files without modifying sources."""

from __future__ import annotations

import argparse
import json
import subprocess
import tempfile
from math import ceil
from pathlib import Path

from openpyxl import load_workbook

IMAGES = {".png", ".jpg", ".jpeg", ".webp"}
MEDIA = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".mp4", ".mov", ".avi", ".mkv"}


def read_text(path: Path) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding), encoding
        except UnicodeDecodeError:
            pass
    return path.read_text(encoding="utf-8", errors="replace"), "utf-8-replace"


def docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    for section in doc.sections:
        blocks.extend(p.text for p in section.header.paragraphs if p.text.strip())
        blocks.extend(p.text for p in section.footer.paragraphs if p.text.strip())
    for number, table in enumerate(doc.tables, 1):
        blocks.append(f"[表格 {number}]")
        for row in table.rows:
            blocks.append("\t".join(cell.text.replace("\n", " ") for cell in row.cells))
    return "\n".join(blocks)


def xlsx_text(path: Path) -> tuple[str, list[str]]:
    formulas = load_workbook(path, data_only=False, read_only=True)
    values = load_workbook(path, data_only=True, read_only=True)
    blocks, sheets = [], []
    for ws_formula in formulas.worksheets:
        if ws_formula.sheet_state != "visible":
            continue
        ws_value = values[ws_formula.title]
        sheets.append(ws_formula.title)
        blocks.append(f"[工作表 {ws_formula.title}]")
        for row_no, (formula_row, value_row) in enumerate(zip(ws_formula.iter_rows(), ws_value.iter_rows()), 1):
            cells = []
            for formula_cell, value_cell in zip(formula_row, value_row):
                value = value_cell.value
                formula = formula_cell.value if formula_cell.data_type == "f" else None
                if value is None and formula is None:
                    continue
                rendered = "" if value is None else str(value)
                if formula:
                    rendered += f" [公式:{formula}]"
                cells.append(f"{formula_cell.coordinate}={rendered}")
            if cells:
                blocks.append(f"行{row_no}\t" + "\t".join(cells))
    return "\n".join(blocks), sheets


def pdf_ocr(path: Path) -> str:
    with tempfile.TemporaryDirectory(prefix="case-material-pdf-") as tmp:
        prefix = Path(tmp) / "page"
        subprocess.run(["pdftoppm", "-png", "-r", "200", str(path), str(prefix)], check=True, capture_output=True)
        pages = []
        for number, image in enumerate(sorted(Path(tmp).glob("page-*.png")), 1):
            result = subprocess.run(
                ["tesseract", str(image), "stdout", "-l", "chi_sim+eng", "--psm", "6"],
                capture_output=True,
                text=True,
            )
            pages.append(f"[第{number}页 OCR]\n{result.stdout.strip()}")
        return "\n\n".join(pages)


def pdf_page_count(path: Path) -> int | None:
    try:
        result = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, check=True)
    except (OSError, subprocess.CalledProcessError):
        return None
    for line in result.stdout.splitlines():
        if line.lower().startswith("pages:"):
            try:
                return int(line.split(":", 1)[1].strip())
            except ValueError:
                return None
    return None


def pdf_text(path: Path) -> tuple[str, str]:
    result = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True)
    text = result.stdout.replace("\f", "\n[分页]\n").strip()
    cjk = sum("\u4e00" <= char <= "\u9fff" for char in text)
    if not text or text.count("■") > 8 or (len(text) > 80 and cjk < 5):
        ocr = pdf_ocr(path)
        return ocr, "PDF逐页OCR（文本层为空或乱码）"
    return text, "PDF文本层"


def image_ocr(path: Path) -> tuple[str, str]:
    from PIL import Image

    with Image.open(path) as image:
        dimensions = f"{image.width}x{image.height}"
    result = subprocess.run(
        ["tesseract", str(path), "stdout", "-l", "chi_sim+eng", "--psm", "6"],
        capture_output=True,
        text=True,
    )
    return result.stdout.strip(), dimensions


def main() -> None:
    parser = argparse.ArgumentParser(description="按 inventory.json 只读提取材料文本")
    parser.add_argument("inventory", type=Path)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    inventory = json.loads(args.inventory.read_text(encoding="utf-8"))
    source = Path(inventory["source_folder"])
    text_dir = args.out_dir / "extracted_text"
    text_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for index, item in enumerate(inventory["files"], 1):
        path = source / item["original_relative_path"]
        suffix, text, method, details, status = path.suffix.lower(), "", "", {}, "已提取"
        try:
            if suffix in MEDIA:
                status, method = "媒体原件已保留，内容依据逐字稿", "不播放、不本地转写"
            elif suffix == ".docx":
                text, method = docx_text(path), "DOCX正文+表格+页眉页脚"
            elif suffix in {".xlsx", ".xlsm"}:
                text, sheets = xlsx_text(path)
                method, details["visible_sheets"] = "XLSX全部可见工作表", sheets
            elif suffix == ".pdf":
                text, method = pdf_text(path)
                if not text:
                    status = "无文本层-需要逐页OCR"
            elif suffix in IMAGES:
                text, dimensions = image_ocr(path)
                method, details["dimensions"] = "Tesseract chi_sim+eng OCR", dimensions
                if not text:
                    status = "OCR失败"
            elif suffix in {".txt", ".md", ".json", ".csv", ".tsv"}:
                text, encoding = read_text(path)
                method, details["encoding"] = "文本读取", encoding
            else:
                status, method = "当前版本不支持解析", "仅登记"
        except Exception as exc:
            status, details["error"] = "解析失败", f"{type(exc).__name__}: {exc}"
        if suffix in {".xlsx", ".xlsm"} and details.get("visible_sheets"):
            details["reading_unit_type"] = "worksheet"
            details["units_expected"] = len(details["visible_sheets"])
            details["source_units"] = list(details["visible_sheets"])
            details["completed_units"] = list(details["visible_sheets"]) if status == "已提取" else []
        elif suffix == ".pdf":
            page_count = pdf_page_count(path)
            if page_count is not None:
                details["reading_unit_type"] = "page"
                details["units_expected"] = page_count
                details["source_units"] = [str(number) for number in range(1, page_count + 1)]
                details["completed_units"] = [str(number) for number in range(1, page_count + 1)] if status == "已提取" else []
        elif suffix in MEDIA:
            details["reading_unit_type"] = "transcript_segment"
            details["units_expected"] = None
            details["source_units"] = []
            details["completed_units"] = []
        else:
            segment_count = max(1, ceil(len(text) / 12000))
            details["reading_unit_type"] = "segment"
            details["units_expected"] = segment_count
            details["source_units"] = [str(number) for number in range(1, segment_count + 1)]
            details["completed_units"] = [str(number) for number in range(1, segment_count + 1)] if status == "已提取" else []
        details["units_completed"] = len(details.get("completed_units", []))
        material_id = f"MAT-{index:04d}"
        text_path = text_dir / f"{material_id}.txt"
        text_path.write_text(text, encoding="utf-8")
        results.append({
            "material_id": material_id,
            "original_relative_path": item["original_relative_path"],
            "method": method,
            "status": status,
            "text_file": text_path.relative_to(args.out_dir).as_posix(),
            "characters": len(text),
            "preview": text[:1200],
            **details,
        })
    output = args.out_dir / "extractions.json"
    output.write_text(json.dumps({"source_folder": str(source), "files": results}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "files": len(results)}, ensure_ascii=False))


if __name__ == "__main__":
    main()

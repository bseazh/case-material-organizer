#!/usr/bin/env python3
"""Build the lawyer-facing Word case report from one confirmed plan."""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from case_naming import report_filename, report_title

NAVY = "17324D"
TEAL = "167D86"
GRAY = "66757C"


def clean(value: object, default: str = "待确认") -> str:
    text = str(value or "").strip()
    text = re.sub(r"(?:EVT|MAT|ENT|PER|ISS)-\d{3,4}(?:\s*[；;、,]\s*)?", "", text)
    text = re.sub(r"\s+", " ", text).strip(" ；;、,")
    return text or default


def split_ids(value: object) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value]
    return re.findall(r"MAT-\d{4}", str(value or ""))


def event_is_main(event: dict) -> bool:
    role = str(event.get("timeline_role") or event.get("timeline_section") or "main").lower()
    return role not in {"background", "背景", "背景信息"}


def event_sort_key(event: dict) -> tuple[bool, str]:
    value = str(event.get("event_time") or "")
    return (re.search(r"(?:19|20)\d{2}", value) is None, value)


def material_names(event: dict, by_id: dict[str, dict]) -> str:
    names = []
    for material_id in split_ids(event.get("all_material_ids") or event.get("material_ids")):
        item = by_id.get(material_id, {})
        name = item.get("proposed_name") or item.get("original_name")
        if name and name not in names:
            names.append(str(name))
    if not names:
        for path in re.split(r"[；;\n]+", str(event.get("archive_paths") or event.get("related_materials") or "")):
            name = Path(path.strip()).name
            if name and name not in names:
                names.append(name)
    return "；".join(names) or "依据材料待确认"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    tc_pr.append(fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "17242A") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(clean(text, ""))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    tr_pr.append(cannot_split)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, "FFFFFF")
        shade(table.rows[0].cells[index], TEAL)
    keep_row_together(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        keep_row_together(table.rows[-1])
        for index, value in enumerate(values):
            set_cell_text(cells[index], value)
            if row_index % 2:
                shade(cells[index], "F4F7F7")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph()


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)


def entity_row(entity: dict) -> list[str]:
    return [
        clean(entity.get("standard_name") or entity.get("name")),
        clean(entity.get("case_roles") or entity.get("role")),
        clean(entity.get("strong_identifier_summary") or entity.get("details") or entity.get("aliases")),
        clean(entity.get("issues"), "无"),
    ]


def write_basic_index(plan: dict, items: list[dict], events: list[dict]) -> None:
    result_folder = Path(str(plan.get("result_folder") or ""))
    basic_folder = result_folder / "002 基础资料"
    if not basic_folder.is_dir():
        return
    event_by_material: dict[str, list[str]] = defaultdict(list)
    for event in events:
        for material_id in split_ids(event.get("all_material_ids") or event.get("material_ids")):
            description = clean(event.get("description"))
            if description not in event_by_material[material_id]:
                event_by_material[material_id].append(description)
    rows = []
    for item in items:
        if item.get("target_category") != "002 基础资料":
            continue
        relative = str(item.get("actual_target_relative_path") or "")
        name = clean(item.get("proposed_name") or item.get("original_name"))
        date_match = re.match(r"^(\d{6}(?:-\d{6})?|\d{4})-", name)
        date_text = date_match.group(1) if date_match else "时间待确认"
        matters = "；".join(event_by_material.get(str(item.get("material_id")), [])) or "对应事项待确认"
        rows.append((date_text == "时间待确认", date_text, name, relative, matters))
    lines = ["# 基础资料索引", "", "| 日期 | 材料名称 | 相对路径 | 对应事项 |", "|---|---|---|---|"]
    for _, date_text, name, relative, matters in sorted(rows):
        safe = lambda value: str(value).replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {safe(date_text)} | {safe(name)} | {safe(relative)} | {safe(matters)} |")
    (basic_folder / "index.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="从已执行归档方案生成案件梳理 Word 报告")
    parser.add_argument("plan", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    plan = json.loads(args.plan.read_text(encoding="utf-8"))
    if not plan.get("confirmed"):
        raise SystemExit("只能根据已确认并执行的归档方案生成案件报告")
    media_check = plan.get("media_check", {})
    if media_check and not media_check.get("ready_for_case_analysis", False):
        raise SystemExit("录音逐字稿检查尚未通过，不能生成案件报告")
    if media_check.get("recording_count", 0) and not plan.get("transcript_mainline_review"):
        raise SystemExit("尚未记录逐字稿候选主线与全量材料反向核查，不能生成案件报告")

    items = plan.get("items", [])
    by_id = {str(item.get("material_id")): item for item in items}
    events = sorted((event for event in plan.get("events", []) if event_is_main(event)), key=event_sort_key)
    summary = plan.get("case_summary", {})
    entities = plan.get("entities", [])

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for style_name, size, color in (("Title", 24, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 12, TEAL)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    case_name = clean(plan.get("case_folder_name"), "案件材料")
    expected_filename = report_filename(plan)
    output = args.out or (Path(str(plan.get("result_folder") or args.plan.parent.parent.parent)) / "整理结果" / expected_filename)
    if output.suffix.lower() != ".docx":
        output = output / expected_filename
    if output.name != expected_filename:
        raise SystemExit(f"报告文件名必须与确认案由一致，应为：{expected_filename}")
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(report_title(plan))
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(case_name)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"依据已归档材料编制 | 生成日期：{date.today().isoformat()} | 材料：{len(items)} 份")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor.from_string(GRAY)

    overview = clean(summary.get("执行摘要") or summary.get("overview"), "")
    if not overview:
        overview = " ".join(clean(summary.get(key), "") for key in ("起因", "过程", "争议", "现状") if summary.get(key))
    add_heading(document, "执行摘要", 1)
    document.add_paragraph(overview or "案件摘要待根据材料补充。")

    add_heading(document, "一、案件主体", 1)
    companies = [entity_row(entity) for entity in entities if str(entity.get("entity_type", "")).lower() in {"企业", "公司", "机构", "company", "organization"}]
    people = [entity_row(entity) for entity in entities if entity_row(entity) not in companies]
    if companies:
        add_heading(document, "1.1 企业及其他机构", 2)
        add_table(document, ["标准名称", "案件角色", "关键信息", "待确认事项"], companies, [4.3, 4.0, 5.0, 4.0])
    if people:
        add_heading(document, "1.2 自然人", 2)
        add_table(document, ["姓名", "案件角色", "关键信息", "待确认事项"], people, [3.3, 4.4, 5.4, 4.2])
    if not entities:
        document.add_paragraph("案件主体待根据材料补充。")

    add_heading(document, "二、案件总结", 1)
    for index, key in enumerate(("起因", "过程", "争议", "现状", "缺口"), 1):
        label = "缺口与待核事项" if key == "缺口" else key
        add_heading(document, f"2.{index} {label}", 2)
        document.add_paragraph(clean(summary.get(key), "待根据现有材料进一步核对。"))

    add_heading(document, "三、关键时间轴", 1)
    timeline_rows = [[
        clean(event.get("event_time"), "时间待确认"),
        clean(event.get("description")),
        clean(event.get("subjects")),
        material_names(event, by_id),
        clean(event.get("conflicts") or event.get("issues"), "无"),
    ] for event in events]
    add_table(document, ["日期", "事件", "涉及主体", "依据材料", "待确认事项"], timeline_rows or [["-", "暂无可确认的主线事件", "-", "-", "待补充材料"]], [2.7, 5.2, 3.5, 4.4, 3.4])

    add_heading(document, "四、文件清单", 1)
    grouped: dict[str, list[dict]] = defaultdict(list)
    for item in items:
        folder = "/".join(part for part in (str(item.get("target_category") or "待分类"), str(item.get("target_subcategory") or "")) if part)
        grouped[folder].append(item)
    for number, (folder, folder_items) in enumerate(sorted(grouped.items()), 1):
        add_heading(document, f"4.{number} {folder}（{len(folder_items)}份）", 2)
        rows = []
        for item in sorted(folder_items, key=lambda row: str(row.get("proposed_name") or row.get("original_name") or "")):
            description = clean(item.get("content_summary") or item.get("summary") or item.get("review_notes"), "材料内容见原件。")
            rows.append([clean(item.get("proposed_name") or item.get("original_name")), description])
        add_table(document, ["文件名", "内容说明"], rows, [7.0, 12.2])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("依据现有材料整理，不构成法律意见；关键事实请以核对原件后的证据为准。")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(GRAY)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    write_basic_index(plan, items, events)
    print(json.dumps({"output": str(output.resolve()), "events": len(events), "materials": len(items)}, ensure_ascii=False))


if __name__ == "__main__":
    main()

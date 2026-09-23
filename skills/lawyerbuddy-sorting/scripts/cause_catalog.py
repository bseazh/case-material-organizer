#!/usr/bin/env python3
"""Search and validate causes of action in the bundled 2025 reference workbook."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

ASSET_DIR = Path(__file__).resolve().parent.parent / "assets"
JSON_CATALOG = ASSET_DIR / "民事案件案由参考表_2025.json"
XLSX_CATALOG = ASSET_DIR / "民事案件案由参考表_2025.xlsx"
DEFAULT_CATALOG = JSON_CATALOG if JSON_CATALOG.is_file() else XLSX_CATALOG
SPLIT = re.compile(r"[；;\n]+")


def text(value: object) -> str:
    return str(value or "").strip()


def load_catalog(path: Path) -> list[dict]:
    if path.suffix.lower() == ".json":
        records = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(records, list) or not all(isinstance(row, dict) for row in records):
            raise ValueError(f"案由参考数据格式无效：{path}")
        return records
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    second_sheet = workbook["二级案由总表"]
    third_sheet = workbook["三级案由清单"]
    second_info: dict[str, dict] = {}
    records: list[dict] = []

    for row in second_sheet.iter_rows(min_row=5, values_only=True):
        level_1, level_2 = text(row[1]), text(row[2])
        if not level_1 or not level_2:
            continue
        info = {
            "level_1": level_1,
            "level_2": level_2,
            "classification": text(row[4]),
            "ai_cues": text(row[5]),
            "boundary": text(row[6]),
        }
        second_info[level_2] = info
        records.append({"level": 2, "name": level_2, "hierarchy": [level_1, level_2], **info})

    level_1_names = sorted({record["level_1"] for record in records})
    for name in level_1_names:
        records.append({"level": 1, "name": name, "hierarchy": [name], "level_1": name})

    for row in third_sheet.iter_rows(min_row=5, values_only=True):
        level_2, level_3, level_4_text = text(row[1]), text(row[3]), text(row[4])
        if not level_2 or not level_3:
            continue
        info = second_info.get(level_2, {"level_1": "", "classification": "", "ai_cues": "", "boundary": ""})
        level_1 = info.get("level_1", "")
        records.append({
            "level": 3,
            "name": level_3,
            "hierarchy": [value for value in (level_1, level_2, level_3) if value],
            "level_1": level_1,
            "level_2": level_2,
            "level_3": level_3,
            "classification": info.get("classification", ""),
            "ai_cues": info.get("ai_cues", ""),
            "boundary": info.get("boundary", ""),
        })
        for level_4 in (part.strip() for part in SPLIT.split(level_4_text) if part.strip()):
            records.append({
                "level": 4,
                "name": level_4,
                "hierarchy": [value for value in (level_1, level_2, level_3, level_4) if value],
                "level_1": level_1,
                "level_2": level_2,
                "level_3": level_3,
                "level_4": level_4,
                "classification": info.get("classification", ""),
                "ai_cues": info.get("ai_cues", ""),
                "boundary": info.get("boundary", ""),
            })
    return records


def compact(record: dict) -> dict:
    return {key: value for key, value in record.items() if value not in ("", None)}


def normalize(value: str) -> str:
    return re.sub(r"[\s，。；、：:（）()《》\-_/]", "", value)


def score(query: str, record: dict) -> int:
    needle = normalize(query)
    name = normalize(record["name"])
    stem = name.removesuffix("纠纷")
    result = 0
    if name and name in needle:
        result += 120
    elif len(stem) >= 2 and stem in needle:
        result += 80
    haystack = normalize(" ".join(text(record.get(key)) for key in ("classification", "ai_cues", "boundary")))
    query_terms = {needle[index:index + 2] for index in range(max(0, len(needle) - 1))}
    result += min(35, sum(1 for term in query_terms if term and term in haystack))
    result += record["level"]
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="检索或校验内置民事案件案由")
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--query", help="用案件关系、诉求和争点检索候选案由")
    group.add_argument("--validate", help="校验一个案由名称并返回层级")
    group.add_argument("--children", help="列出某一级或二级案由下的全部更具体案由")
    parser.add_argument("--limit", type=int, default=8)
    args = parser.parse_args()

    records = load_catalog(args.catalog)
    if args.validate:
        matches = [compact(record) for record in records if record["name"] == args.validate.strip()]
        print(json.dumps({"valid": bool(matches), "matches": matches}, ensure_ascii=False, indent=2))
        raise SystemExit(0 if matches else 2)

    if args.children:
        parent = args.children.strip()
        matches = [
            compact(record) for record in records
            if parent in record.get("hierarchy", []) and record["name"] != parent
        ]
        matches.sort(key=lambda record: (record["level"], record["name"]))
        print(json.dumps({"parent": parent, "matches": matches}, ensure_ascii=False, indent=2))
        raise SystemExit(0 if matches else 2)

    ranked = sorted(
        ((score(args.query or "", record), record) for record in records),
        key=lambda item: (-item[0], -item[1]["level"], item[1]["name"]),
    )
    matches = [compact(record | {"search_score": value}) for value, record in ranked if value > 0][:max(1, args.limit)]
    print(json.dumps({
        "query": args.query,
        "notice": "检索结果仅供缩小候选范围；必须结合请求权、法律关系和排除边界人工复核。",
        "matches": matches,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Render one confirmed case timeline as deterministic, print-ready HTML."""

from __future__ import annotations

import argparse
import html
import json
import os
import re
from collections import Counter
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook


def esc(value: object) -> str:
    return html.escape("" if value is None else str(value))


def split_cn(value: object) -> list[str]:
    return [part.strip() for part in re.split(r"[；;\n]+", str(value or "")) if part.strip()]


def year_of(value: object) -> str:
    match = re.search(r"(?:19|20)\d{2}", str(value or ""))
    return match.group(0) if match else "时间待核"


def date_parts(value: object) -> tuple[str, str]:
    text = str(value or "时间待核")
    year = year_of(text)
    if year == "时间待核":
        return "待核", "日期未明"
    short = text.replace(year + "-", "", 1).replace(year, "", 1).strip(" -")
    short = re.sub(r"至(?:19|20)\d{2}-", "—", short)
    short = short.replace("至", "—").replace("起", " 起").replace("-", ".")
    return year, short or year


def link_for(path: Path, output_dir: Path) -> str:
    relative = Path(os.path.relpath(path, output_dir))
    return "/".join(quote(part) for part in relative.parts)


def event_theme(event: dict[str, object]) -> tuple[str, str, str]:
    """Return a restrained visual category inferred from user-facing event text."""
    text = " ".join(str(event.get(key) or "") for key in ("事件", "相关材料"))
    themes = (
        (("合同", "协议", "委托", "顾问"), "contract", "合同 / 委托", "文"),
        (("付款", "转账", "货款", "工资", "退款", "元"), "payment", "资金往来", "款"),
        (("检测", "抽检", "复检", "测试", "X射线", "分析报告"), "inspection", "检测 / 鉴定", "检"),
        (("聊天", "沟通", "逐字稿", "录音", "会议"), "communication", "沟通记录", "讯"),
        (("工商", "设立", "变更", "备案", "主体"), "entity", "主体信息", "企"),
    )
    for keywords, css_class, label, mark in themes:
        if any(keyword in text for keyword in keywords):
            return css_class, label, mark
    return "general", "事实节点", "事"


def file_kind(name: str) -> str:
    suffix = Path(name).suffix.lstrip(".").upper()
    return suffix[:5] if suffix else "FILE"


def plan_material_names(event: dict, by_id: dict[str, dict]) -> str:
    material_ids = re.findall(r"MAT-\d{4}", str(event.get("all_material_ids") or event.get("material_ids") or ""))
    names = []
    for material_id in material_ids:
        item = by_id.get(material_id, {})
        name = item.get("proposed_name") or item.get("original_name")
        if name and name not in names:
            names.append(str(name))
    if not names:
        names = [Path(value).name for value in split_cn(event.get("archive_paths") or event.get("related_materials"))]
    return "；".join(name for name in names if name)


def load_source(source: Path) -> tuple[list[dict], dict[str, str], int, Path]:
    if source.suffix.lower() == ".json":
        plan = json.loads(source.read_text(encoding="utf-8"))
        if not plan.get("confirmed"):
            raise SystemExit("只能根据已确认并执行的归档方案生成时间轴")
        media_check = plan.get("media_check", {})
        if media_check and not media_check.get("ready_for_case_analysis", False):
            raise SystemExit("录音逐字稿检查尚未通过，不能生成时间轴")
        if media_check.get("recording_count", 0) and not plan.get("transcript_mainline_review"):
            raise SystemExit("尚未记录逐字稿候选主线与全量材料反向核查，不能生成时间轴")
        items = plan.get("items", [])
        by_id = {str(item.get("material_id")): item for item in items}
        events = []
        for event in plan.get("events", []):
            role = str(event.get("timeline_role") or event.get("timeline_section") or "main").lower()
            if role in {"background", "背景", "背景信息"}:
                continue
            events.append({
                "日期": event.get("event_time") or "时间待核",
                "事件": event.get("description") or "事件内容待确认",
                "相关人员/公司": event.get("subjects") or "相关主体待确认",
                "相关材料": plan_material_names(event, by_id),
                "待确认事项": event.get("conflicts") or event.get("issues") or "无",
            })
        overview = {key: str(plan.get("case_summary", {}).get(key) or "") for key in ("起因", "过程", "争议", "现状", "缺口")}
        archive_root = Path(plan.get("result_folder") or source.parent.parent.parent)
        return events, overview, len(plan.get("issues", [])), archive_root

    wb = load_workbook(source, data_only=True, read_only=True)
    ws = wb["案件时间轴"]
    headers = [cell.value for cell in ws[3]]
    events = []
    for values in ws.iter_rows(min_row=4, values_only=True):
        if any(value is not None for value in values):
            events.append(dict(zip(headers, values)))
    overview_ws = wb["案件概览"]
    overview = {
        str(overview_ws.cell(row, 1).value): str(overview_ws.cell(row, 2).value or "")
        for row in range(4, overview_ws.max_row + 1)
        if overview_ws.cell(row, 1).value
    }
    issue_ws = wb["待补材料"]
    issue_count = sum(1 for row in issue_ws.iter_rows(min_row=4, values_only=True) if any(value is not None for value in row))
    archive_root = source.parent.parent if source.parent.name == "整理结果" else source.parent
    return events, overview, issue_count, archive_root


def main() -> None:
    parser = argparse.ArgumentParser(description="从已执行方案JSON生成时间轴HTML（兼容旧版Excel）")
    parser.add_argument("source", type=Path)
    parser.add_argument("--out", type=Path, default=Path("时间轴.html"))
    parser.add_argument("--title", default="案件关键时间轴")
    parser.add_argument("--subtitle", default="已归档材料整理结果 · 事实中性呈现")
    parser.add_argument("--notice", default="仅依据当前已归档材料整理")
    args = parser.parse_args()

    events, overview, issue_count, archive_root = load_source(args.source)
    events.sort(key=lambda event: (str(event.get("日期") or "9999"), str(event.get("事件") or "")))
    years = [year_of(event.get("日期")) for event in events]
    year_counts = Counter(years)
    known_years = {year for year in years if year != "时间待核"}
    material_names = {name for event in events for name in split_cn(event.get("相关材料"))}
    archive_index: dict[str, list[Path]] = {}
    material_folders = sorted(
        path for path in archive_root.iterdir()
        if path.is_dir() and path.name != "整理结果" and not path.name.startswith(".")
    )
    for base in material_folders:
        for path in base.rglob("*"):
            if path.is_file():
                archive_index.setdefault(path.name, []).append(path)

    cards: list[str] = []
    current_year = None
    for event in events:
        event_year = year_of(event.get("日期"))
        if event_year != current_year:
            current_year = event_year
            cards.append(
                f'<div class="year-break"><span>{esc(event_year)}</span>'
                f'<b>{year_counts[event_year]} 个事件</b></div>'
            )
        year, date = date_parts(event.get("日期"))
        names = split_cn(event.get("相关材料"))
        attachments = ""
        print_paths = ""
        if names:
            material_items = []
            for name in names:
                candidates = archive_index.get(name, [])
                label = (
                    f'<a href="{link_for(candidates[0], args.out.parent.resolve())}">{esc(name)}</a>'
                    if len(candidates) == 1 else esc(name)
                )
                material_items.append(
                    f'<li><span class="file-kind">{esc(file_kind(name))}</span>{label}</li>'
                )
            attachments = (
                f'<div class="materials"><div class="materials-head"><b>相关材料</b>'
                f'<span>{len(names)} 份</span></div><ul>{"".join(material_items)}</ul></div>'
            )
            print_paths = f'<p class="print-paths"><b>相关材料</b>{esc("；".join(names))}</p>'
        check = str(event.get("待确认事项") or "无")
        theme, theme_label, theme_mark = event_theme(event)
        check_badge = '<span class="precision">待核</span>' if check not in {"无", "", "无。"} else ""
        cards.append(f'''
<article class="event {theme}">
  <div class="date-block"><span>{esc(year)}</span><strong>{esc(date)}</strong></div>
  <div class="rail"><i>{esc(theme_mark)}</i></div>
  <section>
    <div class="event-top"><span class="tone">{esc(theme_label)}</span>{check_badge}</div>
    <h2>{esc(event.get("事件"))}</h2>
    <dl><div><dt>相关主体</dt><dd>{esc(event.get("相关人员/公司"))}</dd></div></dl>
    {attachments}
    <p class="check"><b>待确认</b>{esc(check)}</p>
    {print_paths}
  </section>
</article>''')

    summary_labels = [("起因", "01"), ("过程", "02"), ("争议", "03"), ("现状", "04"), ("缺口", "05")]
    summary_html = "".join(
        f'<section class="summary-{number}"><span>{number}</span><div><h3>{label}</h3><p>{esc(overview.get(label, "待补充"))}</p></div></section>'
        for label, number in summary_labels
    )
    dated_events = [event for event in events if year_of(event.get("日期")) != "时间待核"]
    first_date = esc(dated_events[0].get("日期") if dated_events else "-")
    last_date = esc(dated_events[-1].get("日期") if dated_events else "-")
    undated_note = f"（另有 {len(events) - len(dated_events)} 项日期待确认）" if len(dated_events) != len(events) else ""
    document = f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{esc(args.title)}</title>
<style>
:root{{--ink:#17242a;--deep:#123f3b;--blue:#205c8f;--teal:#167d72;--gold:#c29132;--red:#b64743;--gray:#66757c;--paper:#fff;--ground:#eef2f1;--line:#bdcbc8}}
*{{box-sizing:border-box}}html{{background:var(--ground)}}body{{margin:0;color:var(--ink);background:var(--ground);font-family:Arial,"PingFang SC","Microsoft YaHei",sans-serif;letter-spacing:0}}
main{{width:min(1180px,calc(100% - 40px));margin:0 auto;padding:48px 0 72px}}
.hero{{position:relative;padding:38px 42px 32px;background:var(--deep);color:#fff;border-top:7px solid var(--gold);overflow:hidden}}
.hero:after{{content:"";position:absolute;right:42px;top:42px;width:90px;height:90px;border:1px solid rgba(255,255,255,.16);box-shadow:18px 18px 0 -1px var(--deep),18px 18px 0 0 rgba(255,255,255,.1)}}
.eyebrow{{margin:0 0 14px;color:#a9d8cf;font-size:12px;font-weight:700}}h1{{max-width:950px;margin:0;font-size:32px;line-height:1.28;letter-spacing:0}}.subtitle{{margin:12px 0 0;color:#d6e7e3;font-size:15px}}
.metrics{{display:grid;grid-template-columns:repeat(4,1fr);margin-top:30px;border-top:1px solid #47716d}}
.metric{{padding:18px 18px 0 0}}.metric b{{display:block;font-size:27px;color:#fff}}.metric span{{color:#bad0dc;font-size:12px}}
.notice{{display:flex;justify-content:space-between;gap:18px;margin-top:24px;padding-top:18px;border-top:1px solid #47716d;color:#d6e7e3;font-size:12px}}
.overview{{margin:28px 0 42px;padding:30px 34px 32px;background:#fff;border:1px solid #d3ddda;border-top:4px solid var(--gold)}}
.section-title{{display:flex;align-items:baseline;justify-content:space-between;margin-bottom:22px}}.section-title h2{{margin:0;font-size:23px;color:var(--deep)}}.section-title p{{margin:0;color:var(--gray);font-size:12px}}
.summary-grid{{display:grid;grid-template-columns:1fr 1fr;gap:0 30px}}.summary-grid section{{display:grid;grid-template-columns:34px 1fr;gap:12px;padding:16px 0;border-top:1px solid #e0e7e5}}.summary-grid section:nth-child(-n+2){{border-top:0;padding-top:0}}.summary-grid span{{display:flex;align-items:center;justify-content:center;width:29px;height:29px;background:#e5f1ee;color:var(--teal);font:700 11px Arial}}.summary-grid h3{{margin:4px 0 7px;font-size:15px;color:var(--deep)}}.summary-grid p{{margin:0;font-size:13px;line-height:1.75;color:#3f5055}}.summary-grid .summary-05{{grid-column:1/-1}}
.legend{{display:flex;flex-wrap:wrap;gap:20px;margin:0 0 22px;padding:0 0 17px;border-bottom:1px solid #c8d4d1;font-size:12px;color:#50616c}}.legend span:before{{content:"";display:inline-block;width:8px;height:8px;margin-right:7px;background:var(--teal)}}.legend .l3:before{{background:var(--blue)}}.legend .l4:before{{background:var(--gold)}}.legend .l5:before{{background:var(--red)}}
.timeline{{position:relative}}.year-break{{display:flex;align-items:center;gap:15px;margin:42px 0 20px 236px;color:var(--deep)}}.year-break:after{{content:"";height:1px;flex:1;background:#afc0bc}}.year-break span{{font-size:28px;font-weight:800}}.year-break b{{padding:5px 8px;background:#dfe9e6;font-size:11px;color:#526760;white-space:nowrap}}
.event{{--accent:var(--teal);display:grid;grid-template-columns:180px 68px 1fr;align-items:stretch;margin-bottom:18px}}
.event.payment{{--accent:#b7792f}}.event.inspection{{--accent:#225f91}}.event.communication{{--accent:#7b5b95}}.event.entity{{--accent:#47756c}}.event.contract{{--accent:#8a6040}}.event.general{{--accent:#68777d}}
.date-block{{align-self:start;padding:14px 16px 13px;border-top:3px solid var(--accent);background:#fff;text-align:right;box-shadow:0 2px 7px rgba(24,52,48,.06)}}.date-block span{{display:block;font-size:11px;color:var(--gray);font-weight:700}}.date-block strong{{display:block;margin-top:4px;color:var(--deep);font-size:23px;line-height:1.1}}
.rail{{position:relative}}.rail:before{{content:"";position:absolute;top:0;bottom:-18px;left:33px;width:2px;background:var(--line)}}.rail i{{position:absolute;z-index:1;top:14px;left:18px;display:flex;align-items:center;justify-content:center;width:31px;height:31px;border:4px solid var(--ground);border-radius:50%;background:var(--accent);box-shadow:0 0 0 2px var(--line);color:#fff;font-style:normal;font-size:11px;font-weight:700}}
.event section{{padding:22px 26px 20px;background:#fff;border:1px solid #d4dedb;border-left:5px solid var(--accent);box-shadow:0 3px 10px rgba(24,52,48,.055)}}.event-top{{display:flex;align-items:center;gap:8px;font-size:11px}}.tone,.precision{{padding:4px 7px;background:#edf3f1;color:#425761;font-weight:700}}.precision{{background:#f8eceb;color:#963d3a}}
.event h2{{margin:10px 0 14px;color:var(--ink);font-size:19px;line-height:1.55}}
dl{{margin:0 0 13px}}dl div{{display:grid;grid-template-columns:76px 1fr;gap:8px}}dt,.check>b{{color:#687883;font-size:11px;font-weight:700}}dd{{margin:0;font-size:12px;line-height:1.6}}
.materials{{display:grid;grid-template-columns:76px 1fr;gap:8px;margin:10px 0 0}}.materials-head{{font-size:11px;color:#687883;font-weight:700}}.materials-head span{{display:block;margin-top:3px;color:#89959a;font-size:10px;font-weight:400}}.materials ul{{display:flex;flex-wrap:wrap;gap:6px;margin:0;padding:0;list-style:none}}.materials li{{display:flex;align-items:center;max-width:100%;border:1px solid #d7dfdc;background:#f7f9f8;color:#435359;font-size:10px}}.materials a,.materials li{{overflow-wrap:anywhere}}.materials a{{padding:5px 7px 5px 0;color:#285b6c;text-decoration:none}}.file-kind{{align-self:stretch;display:flex;align-items:center;margin-right:7px;padding:4px 5px;background:#e4ece9;color:#526b64;font:700 8px Arial}}
.check{{display:grid;grid-template-columns:76px 1fr;gap:8px;margin:12px 0 0;padding-top:11px;border-top:1px solid #ead9d9;color:#8e3030;font-size:12px;line-height:1.6}}.check b{{color:var(--red)}}
a:hover{{text-decoration:underline}}.print-paths{{display:none}}
.footer{{margin:34px 0 0 248px;padding:18px 0;border-top:1px solid #b9c6cc;color:#61717b;font-size:11px;line-height:1.7}}
@media(max-width:760px){{main{{width:calc(100% - 16px);padding-top:8px}}.hero{{padding:25px 18px}}.hero:after{{display:none}}h1{{max-width:none;font-size:22px}}.metrics{{grid-template-columns:1fr 1fr}}.notice{{display:block}}.overview{{padding:24px 20px}}.summary-grid{{grid-template-columns:1fr}}.summary-grid section,.summary-grid section:nth-child(-n+2){{grid-column:auto;padding:14px 0;border-top:1px solid #e1e6e9}}.summary-grid section:first-child{{padding-top:0;border-top:0}}.event{{grid-template-columns:58px 28px 1fr;margin-bottom:14px}}.year-break{{margin:30px 0 14px 86px}}.year-break span{{font-size:23px}}.date-block{{padding:9px 5px}}.date-block strong{{font-size:15px}}.rail:before{{left:13px;bottom:-14px}}.rail i{{top:11px;left:2px;width:24px;height:24px;border-width:3px;font-size:9px}}.event section{{padding:16px 14px 14px}}.event h2{{font-size:16px;line-height:1.5}}dl div,.check{{grid-template-columns:58px 1fr}}.materials{{grid-template-columns:1fr}}.materials ul{{display:grid;grid-template-columns:1fr}}.materials li{{width:100%}}.footer{{margin-left:86px}}}}
@media print{{@page{{size:A4;margin:13mm 11mm}}html,body{{background:#fff}}main{{width:100%;padding:0}}.hero{{padding:20px 24px 18px;-webkit-print-color-adjust:exact;print-color-adjust:exact}}.hero:after{{display:none}}h1{{font-size:25px}}.metrics{{margin-top:14px}}.metric{{padding-top:10px}}.metric b{{font-size:20px}}.notice{{display:grid;grid-template-columns:1fr auto;column-gap:28px;font-size:9px}}.notice span:last-child{{text-align:right}}.overview{{margin:14px 0 20px;padding:16px 20px}}.summary-grid p{{font-size:9px;line-height:1.45}}.timeline{{padding-top:1px}}.year-break{{margin-top:18px;margin-bottom:10px;break-after:avoid}}.event{{grid-template-columns:108px 36px 1fr;margin-bottom:9px;break-inside:avoid;page-break-inside:avoid}}.date-block{{padding:8px}}.date-block strong{{font-size:14px}}.rail:before{{left:17px;bottom:-9px}}.rail i{{top:8px;left:7px;width:21px;height:21px;border-width:3px;font-size:8px}}.event section{{padding:12px 15px 10px}}.event h2{{font-size:13px;margin:6px 0 8px}}dl{{margin-bottom:6px}}dd,.check{{font-size:9px}}.materials{{display:none}}.print-paths{{display:grid;grid-template-columns:76px 1fr;gap:8px;margin:5px 0 0;font-size:8px;line-height:1.45;color:#53646f;overflow-wrap:anywhere}}.print-paths b{{color:#687883;font-size:8px}}.footer{{display:none}}}}
</style></head>
<body><main>
<header class="hero"><p class="eyebrow">LEGAL AI · MATERIAL-BASED TIMELINE</p><h1>{esc(args.title)}</h1><p class="subtitle">{esc(args.subtitle)}</p>
<div class="metrics"><div class="metric"><b>{len(events)}</b><span>主线事件</span></div><div class="metric"><b>{len(material_names)}</b><span>关联材料</span></div><div class="metric"><b>{len(known_years)}</b><span>涉及年度</span></div><div class="metric"><b>{issue_count}</b><span>待补事项</span></div></div>
<div class="notice"><span>时间跨度：{first_date} — {last_date}{esc(undated_note)}</span><span>{esc(args.notice)}</span></div></header>
<section class="overview"><div class="section-title"><h2>案件概览</h2><p>起因—过程—争议—现状—缺口</p></div><div class="summary-grid">{summary_html}</div></section>
<div class="legend"><span>时间与事件</span><span class="l3">相关人员/公司</span><span class="l4">相关材料</span><span class="l5">待确认事项</span></div>
<div class="timeline">{"".join(cards)}</div>
<footer class="footer">本时间轴仅依据已归档材料整理，不构成事实认定或法律结论。音视频内容、文字识别不清之处，以及主体、日期、金额不一致之处，均应回查原件。</footer>
</main></body></html>'''
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(document, encoding="utf-8")
    print(args.out.resolve())


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Generate technical indexes plus a plain-language, user-facing workbook."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build_tree import OUTPUT_FOLDER, parse_totals, render_tree_markdown, tree_lines
from case_naming import directory_structure

SKILL = Path(__file__).resolve().parents[1]
TEMPLATE = SKILL / "assets" / "案件材料汇总模板.xlsx"


def clear_examples(ws) -> None:
    if ws.max_row >= 4:
        ws.delete_rows(4, ws.max_row - 3)


def material_ids(value: object) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value if re.fullmatch(r"MAT-\d{4}", str(item))]
    return re.findall(r"MAT-\d{4}", str(value or ""))


def clean_user_text(value: object) -> str:
    """Remove internal IDs from prose shown to non-technical users."""
    text = str(value or "").strip()
    text = re.sub(r"(?:EVT|MAT|ENT|PER|ISS)-\d{3,4}(?:\s*[、；;,至-]\s*(?:(?:EVT|MAT|ENT|PER|ISS)-)?\d{3,4})*", "", text)
    text = re.sub(r"（\s*[、；;,。\s]*）", "", text)
    text = re.sub(r"\(\s*[、；;,。\s]*\)", "", text)
    text = re.sub(r"[；;][ \t]*[；;]", "；", text)
    text = re.sub(r"[、,][ \t]*[、,]", "、", text)
    for old, new in (
        ("OCR置信较低", "文字识别清晰度较低"),
        ("OCR低置信", "文字识别清晰度较低"),
        ("解析失败", "暂未识别内容"),
        ("未解析", "需人工查看"),
        ("已解析", "已整理"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ；;、,")


def readable_date(name: str) -> str:
    match = re.match(r"^(\d{6})(?:-(\d{6}))?-", name)
    if match:
        def expand(token: str) -> str:
            return f"20{token[:2]}-{token[2:4]}-{token[4:6]}"
        first = expand(match.group(1))
        return f"{first} 至 {expand(match.group(2))}" if match.group(2) else first
    match = re.match(r"^(\d{4})-", name)
    if match:
        token = match.group(1)
        return f"20{token[:2]}-{token[2:4]}"
    return "时间待确认"


def event_sort_key(event: dict) -> tuple[bool, str]:
    value = str(event.get("event_time") or "")
    return (re.search(r"(?:19|20)\d{2}", value) is None, value)


def event_section(event: dict) -> str:
    """Separate the dispute chain from corporate-history background."""
    explicit = str(event.get("timeline_role") or event.get("timeline_section") or "").strip().lower()
    if explicit in {"background", "背景", "背景信息"}:
        return "background"
    if explicit in {"main", "主线", "案件主线"}:
        return "main"
    description = str(event.get("description") or "")
    background_phrases = ("工商档案", "设立登记", "变更或备案", "工商变更", "历史沿革")
    return "background" if any(phrase in description for phrase in background_phrases) else "main"


def material_type(item: dict) -> str:
    name = item.get("proposed_name", "")
    category = item.get("target_category", "")
    if category == "001 主体信息":
        return "主体/工商材料"
    if category == "003 委托材料":
        return "委托材料"
    if category == "004 类案及法律检索":
        return "法律检索"
    if category == "005 法律文书":
        return "法律文书"
    for keyword, label in (
        ("合同", "合同"), ("协议", "协议"), ("转账", "转账记录"),
        ("银行", "银行记录"), ("聊天", "聊天记录"), ("检测报告", "检测报告"),
        ("录音", "录音"), ("逐字稿", "逐字稿"), ("照片", "图片/照片"),
    ):
        if keyword in name:
            return label
    return {
        ".pdf": "PDF文档", ".doc": "Word文档", ".docx": "Word文档",
        ".xls": "Excel表格", ".xlsx": "Excel表格", ".csv": "表格数据",
        ".jpg": "图片", ".jpeg": "图片", ".png": "图片", ".txt": "文本材料",
        ".m4a": "录音", ".mp3": "录音", ".wav": "录音", ".mp4": "视频", ".mov": "视频",
    }.get(str(item.get("extension", "")).lower(), "其他材料")


def friendly_status(item: dict, has_transcript: bool = False) -> str:
    status = str(item.get("parse_status", ""))
    if status == "当前版本不处理":
        return "已有对应文字材料" if has_transcript else "需人工查看，待提供逐字稿"
    if "OCR低置信" in status:
        return "需核对原件"
    if "失败" in status or "错误" in status:
        return "暂未识别内容"
    if status:
        return "已整理"
    return "待确认"


def transcript_matches(items: list[dict]) -> dict[str, dict]:
    text_extensions = {".doc", ".docx", ".txt", ".pdf"}
    text_items = [x for x in items if str(x.get("extension", "")).lower() in text_extensions]
    matches: dict[str, dict] = {}
    for media in (x for x in items if x.get("parse_status") == "当前版本不处理"):
        tokens = set(re.findall(r"\d{8,14}", media.get("original_name", "")))
        candidates = [item for item in text_items if tokens and tokens.intersection(re.findall(r"\d{8,14}", item.get("original_name", "")))]
        if len(candidates) == 1:
            matches[media["material_id"]] = candidates[0]
    return matches


def main() -> None:
    parser = argparse.ArgumentParser(description="根据已执行归档方案生成索引和用户版汇总表")
    parser.add_argument("applied_plan", type=Path)
    args = parser.parse_args()
    plan = json.loads(args.applied_plan.read_text(encoding="utf-8"))
    if not plan.get("confirmed") or not plan.get("result_folder"):
        raise SystemExit("需要 apply_plan.py 生成的已确认方案")

    result = Path(plan["result_folder"])
    items = plan["items"]
    summary = plan.get("case_summary") or {}
    item_by_id = {item["material_id"]: item for item in items}
    events = sorted(plan.get("events", []), key=event_sort_key)
    main_events = [event for event in events if event_section(event) == "main"]
    background_events = [event for event in events if event_section(event) == "background"]
    media_matches = transcript_matches(items)
    media_rows = [x for x in items if x.get("parse_status") == "当前版本不处理"]
    parsed_count, unparsed_count = parse_totals(items)
    output_dir = result / OUTPUT_FOLDER
    output_dir.mkdir(parents=True, exist_ok=True)

    event_by_material: dict[str, list[str]] = {}
    for event in events:
        description = clean_user_text(event.get("description", ""))
        for material_id in material_ids(event.get("all_material_ids")):
            event_by_material.setdefault(material_id, []).append(description)

    if "002 基础资料" in directory_structure(plan):
        basics = sorted((x for x in items if x["target_category"] == "002 基础资料"), key=lambda x: (x["proposed_name"].startswith("时间待核"), x["proposed_name"]))
        md = ["# 基础资料索引", "", f"共 {len(basics)} 个文件。", "", "| 日期 | 材料名称 | 相对路径 | 对应事项 |", "|---|---|---|---|"]
        for item in basics:
            linked_events = "<br>".join(event_by_material.get(item["material_id"], ["未形成独立时间轴事项，仅作材料索引"]))
            md.append(f"| {readable_date(item['proposed_name'])} | {Path(item['actual_target_relative_path']).name} | {item['actual_target_relative_path']} | {linked_events} |")
        (result / "002 基础资料" / "index.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    output = output_dir / "案件材料汇总.xlsx"
    shutil.copy2(TEMPLATE, output)
    wb = load_workbook(output)
    for ws in wb.worksheets:
        clear_examples(ws)

    overview = wb["案件概览"]
    for row, key in enumerate(("起因", "过程", "争议", "现状", "缺口"), 4):
        overview.cell(row, 1, key)
        overview.cell(row, 2, clean_user_text(summary.get(key, "待补充")))
    if background_events:
        start_row = 10
        overview.cell(start_row, 1, "背景信息")
        overview.cell(start_row, 2, "以下事项用于了解主体历史，不纳入案件主时间轴。")
        for row, event in enumerate(background_events, start_row + 1):
            overview.cell(row, 1, event.get("event_time", "时间待确认"))
            overview.cell(row, 2, clean_user_text(event.get("description", "")))

    materials = wb["材料清单"]
    for row, item in enumerate(items, 4):
        descriptions = list(dict.fromkeys(event_by_material.get(item["material_id"], [])))
        content = "；".join(descriptions) if descriptions else clean_user_text(item.get("review_notes", "未纳入时间轴，仅作材料索引。"))
        archive_folder = Path(item["actual_target_relative_path"]).parent.as_posix()
        values = [readable_date(item["proposed_name"]), Path(item["actual_target_relative_path"]).name, material_type(item), content, archive_folder, friendly_status(item, item["material_id"] in media_matches)]
        for col, value in enumerate(values, 1):
            materials.cell(row, col, value)

    parties = wb["当事人信息"]
    for row, entity in enumerate(plan.get("entities", []), 4):
        details = []
        if entity.get("entity_type"):
            details.append(f"类型：{entity['entity_type']}")
        aliases = str(entity.get("aliases", "")).strip()
        if aliases and aliases not in {"无", "-"}:
            details.append(f"其他称呼：{aliases}")
        if entity.get("confidence") and entity.get("confidence") != "高":
            details.append(f"识别状态：{entity['confidence']}置信")
        values = [entity.get("standard_name", ""), entity.get("case_roles", ""), "；".join(details) or "信息来自归档材料。", entity.get("issues", "")]
        for col, value in enumerate(values, 1):
            parties.cell(row, col, clean_user_text(value))

    timeline = wb["案件时间轴"]
    for row, event in enumerate(main_events, 4):
        linked_names = []
        for material_id in material_ids(event.get("all_material_ids")):
            item = item_by_id.get(material_id)
            if item:
                linked_names.append(Path(item["actual_target_relative_path"]).name)
        values = [event.get("event_time", "时间待确认"), clean_user_text(event.get("description", "")), clean_user_text(event.get("subjects", "")), "；".join(linked_names), clean_user_text(event.get("conflicts", "无"))]
        for col, value in enumerate(values, 1):
            timeline.cell(row, col, value)

    issues = wb["待补材料"]
    for row, issue in enumerate(plan.get("issues", []), 4):
        values = [issue.get("issue", ""), issue.get("impact", ""), issue.get("recommended_action", ""), issue.get("status", "")]
        for col, value in enumerate(values, 1):
            issues.cell(row, col, clean_user_text(value))

    if media_rows:
        media_ws = wb["音视频材料"]
        for row, item in enumerate(media_rows, 4):
            transcript = media_matches.get(item["material_id"])
            transcript_name = Path(transcript["actual_target_relative_path"]).name if transcript else "未找到"
            handling = "已有对应文字材料" if transcript else "待补充逐字稿"
            note = "仅整理文字材料；音视频内容仍需人工确认。" if transcript else "未纳入事实整理，需提供逐字稿。"
            for col, value in enumerate([Path(item["actual_target_relative_path"]).name, transcript_name, handling, note], 1):
                media_ws.cell(row, col, value)
    else:
        wb.remove(wb["音视频材料"])

    preferred_order = ["案件概览", "案件时间轴", "当事人信息", "材料清单", "待补材料", "音视频材料"]
    wb._sheets = [wb[name] for name in preferred_order if name in wb.sheetnames]

    for ws in wb.worksheets:
        for data_row in ws.iter_rows(min_row=4):
            for cell in data_row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.auto_filter.ref = f"A3:{get_column_letter(ws.max_column)}{max(ws.max_row, 3)}"
        ws.freeze_panes = "A4"
    if background_events:
        for cell in overview[10]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="167D86")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    wb.save(output)
    technical_dir = output_dir / "技术资料"
    technical_dir.mkdir(parents=True, exist_ok=True)
    tree_output = technical_dir / "归档结果目录.md"
    index_output = output_dir / "材料统计与目录.txt"
    (output_dir / "index.txt").unlink(missing_ok=True)
    (output_dir / "归档结果目录.md").unlink(missing_ok=True)
    index_output.touch(exist_ok=True)
    tree_output.touch(exist_ok=True)
    tree_output.write_text(render_tree_markdown(plan, "result"), encoding="utf-8")
    lines = [
        "案件材料整理结果", "", "一、材料统计",
        f"材料总数：{len(items)}", f"已整理材料：{parsed_count}", f"需人工查看：{unparsed_count}",
        "", "二、目录树", *tree_lines(plan, "result"), "",
    ]
    index_output.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"result": str(result), "excel": str(output), "tree": str(tree_output), "index": str(index_output), "files": len(items)}, ensure_ascii=False))


if __name__ == "__main__":
    main()

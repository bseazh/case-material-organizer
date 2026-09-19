#!/usr/bin/env python3
"""Create the standard workbook template bundled with this Skill."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parents[1] / "assets" / "案件材料汇总模板.xlsx"
NAVY, TEAL, PALE, BLUE = "17324D", "167D86", "E8F3F4", "0000FF"
SHEETS = {
    "案件概览": ["项目", "核心内容"],
    "材料清单": ["日期", "材料名称", "材料类型", "主要内容", "所在文件夹", "处理状态"],
    "当事人信息": ["姓名/公司", "身份或角色", "相关说明", "待确认事项"],
    "音视频材料": ["文件名称", "对应逐字稿", "处理情况", "备注"],
    "案件时间轴": ["日期", "事件", "相关人员/公司", "相关材料", "待确认事项"],
    "问题与待补材料": ["问题", "可能影响", "建议补充", "处理状态"],
}


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in SHEETS.items():
        ws = wb.create_sheet(name)
        ws["A1"] = name
        ws["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws["A2"] = "用户阅读版：仅展示理解案件所需信息；不确定内容统一标记为‘待确认’。"
        ws["A2"].font = Font(name="Arial", size=10, color=BLUE)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        for col, header in enumerate(headers, 1):
            cell = ws.cell(3, col, header)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=TEAL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        example = ["示例（生成时删除）"] + ["示例值"] * (len(headers) - 1)
        if name == "案件概览":
            example = ["起因", "材料显示双方因合同履行问题发生争议。"]
        elif name == "材料清单":
            example = ["2025-03-03", "首付款回单.pdf", "银行记录", "记载首付款支付情况。", "002 基础资料", "已读取"]
        elif name == "案件时间轴":
            example = ["2025-03-03", "买方支付首付款。", "甲公司；乙公司", "首付款回单.pdf", "需核对银行原始流水。"]
        for col, value in enumerate(example, 1):
            cell = ws.cell(4, col, value)
            cell.font = Font(name="Arial", color=BLUE)
            cell.fill = PatternFill("solid", fgColor=PALE)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{ws.cell(3, len(headers)).coordinate}"
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[3].height = 32
        widths = {
            "案件概览": [16, 90],
            "材料清单": [15, 34, 18, 60, 22, 20],
            "当事人信息": [30, 32, 52, 50],
            "音视频材料": [38, 38, 28, 60],
            "案件时间轴": [18, 70, 38, 60, 60],
            "问题与待补材料": [60, 36, 55, 18],
        }[name]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(3, col).column_letter].width = width
        ws.sheet_view.showGridLines = False
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
